import os
import pandas as pd
from datetime import datetime

from sqlalchemy import text

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.core.logger import logger


def generate_attendance_report(db):

    try:

        logger.info(
            "Attendance report generation started."
        )

        # Generate daily report file name
        report_date = datetime.now().strftime(
            "%d-%m-%Y"
        )

        output_file = os.path.join(
            "files",
            "reports",
            f"Attendance_Report_{report_date}.xlsx"
        )

        logger.info(
            f"Report output file: {output_file}"
        )

        # Fetch attendance data
        query = text("""
            SELECT
                a.EmployeeId,
                e.EmployeeName,
                e.Department,
                a.AttendanceDate,
                a.Status,
                a.LoginTime,
                a.LogoutTime,
                a.SourceFile
            FROM Attendance a
            INNER JOIN Employees e
                ON a.EmployeeId = e.EmployeeId
            ORDER BY
                a.AttendanceDate,
                a.EmployeeId
        """)

        result = db.execute(query)

        data = result.mappings().all()

        df = pd.DataFrame(data)

        # Create report directory
        report_directory = os.path.dirname(
            output_file
        )

        if report_directory:

            os.makedirs(
                report_directory,
                exist_ok=True
            )

        # Generate Excel file
        df.to_excel(
            output_file,
            index=False,
            sheet_name="Attendance Report"
        )

        logger.info(
            "Excel report created successfully."
        )

        # Open generated Excel
        workbook = load_workbook(
            output_file
        )

        worksheet = workbook[
            "Attendance Report"
        ]

        # Header formatting
        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        header_font = Font(
            bold=True,
            color="FFFFFF"
        )

        for cell in worksheet[1]:

            cell.fill = header_fill

            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center"
            )

        # Auto column width
        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = max_length + 3

        # Date formatting
        for cell in worksheet["D"][1:]:

            if cell.value:

                cell.number_format = (
                    "dd-mm-yyyy"
                )

        # Login time formatting
        for cell in worksheet["F"][1:]:

            if cell.value:

                cell.number_format = (
                    "hh:mm:ss"
                )

        # Logout time formatting
        for cell in worksheet["G"][1:]:

            if cell.value:

                cell.number_format = (
                    "hh:mm:ss"
                )

        # Status formatting
        for row in worksheet.iter_rows(
            min_row=2,
            min_col=5,
            max_col=5
        ):

            status_cell = row[0]

            if status_cell.value == "Present":

                status_cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="C6EFCE"
                )

                status_cell.font = Font(
                    color="006100",
                    bold=True
                )

            elif status_cell.value == "Absent":

                status_cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="FFC7CE"
                )

                status_cell.font = Font(
                    color="9C0006",
                    bold=True
                )

        # Center important columns
        for row in worksheet.iter_rows(
            min_row=2
        ):

            # Employee ID
            row[0].alignment = Alignment(
                horizontal="center"
            )

            # Attendance Date
            row[3].alignment = Alignment(
                horizontal="center"
            )

            # Status
            row[4].alignment = Alignment(
                horizontal="center"
            )

            # Login Time
            row[5].alignment = Alignment(
                horizontal="center"
            )

            # Logout Time
            row[6].alignment = Alignment(
                horizontal="center"
            )

        # Freeze header
        worksheet.freeze_panes = "A2"

        # Enable filter
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        # Save final formatted report
        workbook.save(
            output_file
        )

        logger.info(
            "Attendance report generated successfully: "
            f"{output_file}"
        )

        return output_file

    except Exception as ex:

        logger.exception(
            f"Attendance report generation failed. {ex}"
        )

        raise